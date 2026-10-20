#!/usr/bin/env python3
"""
Веб-додаток «ATB & Сільпо PRO: Аналітика, Знижки & Порівняння цін».
Об'єднує парсери інтернет-магазинів «АТБ» та «Сільпо» з паралельним пошуком,
порівнянням цін, 3 режимами відображення магазинів, 3D картками, Badoo-свайпінгом,
розумним спліт-кошиком, сканером штрихкодів, аналізатором чеків та радаром рекордних знижок.
"""

import os
import io
import json
import csv
import time
import re
import urllib.parse
import webbrowser
import threading
from concurrent.futures import ThreadPoolExecutor
from flask import Flask, render_template, request, jsonify, send_file, redirect

import atb_api
import silpo_api
from silpo_mcp import SilpoMCPClient, TOKEN_FILE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE_DIR, "database.json")
ENV_FILE = os.path.join(BASE_DIR, ".env")

# Завантаження змінних оточення з .env
if os.path.exists(ENV_FILE):
    try:
        with open(ENV_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    os.environ.setdefault(k.strip(), v.strip())
    except Exception as e:
        print(f"Помилка завантаження .env: {e}")

app = Flask(__name__, template_folder="templates", static_folder="static")
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.jinja_env.auto_reload = True

mcp_client = SilpoMCPClient()
_oauth_verifier_cache = {}
_db_lock = threading.Lock()


# ═══════════════════════════════════════════
#   РОБОТА З БАЗОЮ ДАНИХ (DATABASE.JSON)
# ═══════════════════════════════════════════
def get_default_db():
    return {
        "wishlist": [],
        "basket": [],
        "settings": {
            "has_atb_card": True,
            "has_silpo_card": True,
            "weekend_discount_active": True
        },
        "price_history": {},
        "receipts": []
    }


def load_db():
    with _db_lock:
        if not os.path.exists(DB_FILE):
            data = get_default_db()
            save_db(data)
            return data
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return get_default_db()


def save_db(data):
    with _db_lock:
        try:
            with open(DB_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            print(f"Помилка запису в database.json: {e}")
            return False


# ═══════════════════════════════════════════
#   ГОЛОВНІ ВЕБ-СТОРІНКИ ТА ФІЛІАЛИ
# ═══════════════════════════════════════════
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/branches")
def get_branches():
    """Повертає список доступних міст та перевірених супермаркетів Сільпо."""
    branches_list = []
    for name, data in silpo_api.CITIES_BRANCHES.items():
        branches_list.append({
            "name": name,
            "city": data["city"],
            "address": data.get("address", ""),
            "branchId": data["branchId"]
        })
    return jsonify({
        "branches": branches_list,
        "defaultBranchId": silpo_api.DEFAULT_BRANCH_ID
    })


# ═══════════════════════════════════════════
#   ПОШУК ТОВАРІВ (АТБ + СІЛЬПО)
# ═══════════════════════════════════════════
# Кеш для акційних товарів (5 хвилин)
_PROMO_CACHE = {
    "timestamp": 0,
    "branch_id": "",
    "atb": [],
    "silpo": []
}


def get_cached_promos(branch_id: str, limit: int = 60):
    now = time.time()
    if (
        now - _PROMO_CACHE["timestamp"] < 300
        and _PROMO_CACHE["branch_id"] == branch_id
        and _PROMO_CACHE["atb"]
        and _PROMO_CACHE["silpo"]
    ):
        return list(_PROMO_CACHE["atb"]), list(_PROMO_CACHE["silpo"])

    atb_items = []
    silpo_items = []
    with ThreadPoolExecutor(max_workers=2) as executor:
        f_atb = executor.submit(atb_api.get_atb_promos, (1, 2, 3), True, 12)
        f_silpo = executor.submit(silpo_api.get_silpo_promotions, branch_id, limit, 12)
        try:
            atb_items = f_atb.result() or []
        except Exception as e:
            print(f"Помилка завантаження акцій АТБ: {e}")
        try:
            silpo_items = f_silpo.result() or []
        except Exception as e:
            print(f"Помилка завантаження акцій Сільпо: {e}")

    _PROMO_CACHE["timestamp"] = now
    _PROMO_CACHE["branch_id"] = branch_id
    _PROMO_CACHE["atb"] = atb_items
    _PROMO_CACHE["silpo"] = silpo_items
    return list(atb_items), list(silpo_items)


# ═══════════════════════════════════════════
#   РОЗУМНІ КАТЕГОРІЇ (SMART CATEGORIES)
# ═══════════════════════════════════════════
SMART_CATEGORIES = {
    "напої": {
        "title": "🥤 Напої, соки, води та енергетики",
        "sub_queries": [
            ["напій", "сік", "енергетик"],
            ["нектар", "вода", "лимонад"],
            ["кола", "квас", "чай"]
        ],
        "aliases": ["напої", "напій", "drinks", "напитки", "попити", "газировка", "соки", "нектари", "енергетики"]
    },
    "молочка": {
        "title": "🥛 Молочні продукти",
        "sub_queries": [
            ["молоко", "сир", "йогурт"],
            ["кефір", "сметана", "масло"],
            ["творог", "вершки", "ряжанка"]
        ],
        "aliases": ["молочка", "молочні", "молочні продукти", "молочное"]
    },
    "м'ясо": {
        "title": "🥩 М'ясо та ковбаси",
        "sub_queries": [
            ["м'ясо", "ковбаса", "сосиски"],
            ["курка", "філе", "фарш"],
            ["свинина", "балик", "сардельки"]
        ],
        "aliases": ["м'ясо", "мясо", "ковбаси", "ковбаса", "сосиски", "м'ясні"]
    },
    "солодощі": {
        "title": "🍫 Солодощі та десерти",
        "sub_queries": [
            ["шоколад", "цукерки", "печиво"],
            ["вафлі", "батончик", "торт"],
            ["круасан", "мармелад", "халва"]
        ],
        "aliases": ["солодощі", "солодке", "десерти", "цукерки", "печиво", "сладости"]
    },
    "снеки": {
        "title": "🥔 Снеки та чипси",
        "sub_queries": [
            ["чипси", "сухарики", "горішки"],
            ["снеки", "попкорн", "арахіс"]
        ],
        "aliases": ["снеки", "чипси", "чіпси", "снек", "хрустики"]
    },
    "риба": {
        "title": "🐟 Риба та морепродукти",
        "sub_queries": [
            ["риба", "лосось", "скумбрія"],
            ["оселедець", "креветки", "ікра"],
            ["тунець", "хек", "мідії"]
        ],
        "aliases": ["риба", "морепродукти", "рыба", "рибне"]
    },
    "овочі": {
        "title": "🥦 Свіжі овочі",
        "sub_queries": [
            ["помідори", "огірки", "картопля"],
            ["капуста", "морква", "цибуля"],
            ["перець", "буряк", "зелень"]
        ],
        "aliases": ["овочі", "овощи", "зелень"]
    },
    "фрукти": {
        "title": "🍎 Свіжі фрукти",
        "sub_queries": [
            ["яблука", "банани", "апельсини"],
            ["лимони", "мандарини", "виноград"],
            ["груші", "ківі", "ананас"]
        ],
        "aliases": ["фрукти", "ягоди", "фрукты"]
    },
    "бакалія": {
        "title": "🌾 Бакалія, крупи та макарони",
        "sub_queries": [
            ["макарони", "гречка", "рис"],
            ["крупа", "борошно", "олія"],
            ["цукор", "сіль", "вівсянка"]
        ],
        "aliases": ["бакалія", "крупи", "макарони", "бакалея"]
    },
    "кава": {
        "title": "☕ Кава та чай",
        "sub_queries": [
            ["кава", "чай"]
        ],
        "aliases": ["кава та чай", "гарячі напої", "чаї"]
    },
    "хімія": {
        "title": "🧼 Побутова хімія",
        "sub_queries": [
            ["порошок", "мило", "шампунь"],
            ["миючий засіб", "зубна паста", "папір"]
        ],
        "aliases": ["побутова хімія", "хімія", "гігієна", "догляд"]
    }
}

def find_smart_category(raw_q: str):
    q = raw_q.lower().strip()
    for cat_id, cat_info in SMART_CATEGORIES.items():
        if q == cat_id or q in cat_info["aliases"]:
            return cat_id, cat_info
        for alias in cat_info["aliases"]:
            if alias in q and len(alias) >= 4:
                return cat_id, cat_info
    return None, None


# ═══════════════════════════════════════════
#   ПОШУК ТОВАРІВ (АТБ + СІЛЬПО)
# ═══════════════════════════════════════════
@app.route("/api/search")
def search():
    """
    Уніфікований пошук товарів з підтримкою вибору магазину:
    - Якщо пошуковий запит порожній (стартова сторінка) або q='акції' — повертає гарячі акційні товари АТБ & Сільпо
    - store=all: паралельний запит до АТБ та Сільпо через ThreadPoolExecutor
    - store=atb: тільки АТБ
    - store=silpo: тільки Сільпо
    """
    raw_query = request.args.get("q", "").strip()
    store = request.args.get("store", "all").strip().lower()
    branch_id = request.args.get("branch_id", silpo_api.DEFAULT_BRANCH_ID).strip()
    limit = int(request.args.get("limit", 48))
    page = int(request.args.get("page", 1))
    sort = request.args.get("sort", "default").strip()
    only_promo = request.args.get("promo", "0") in ("1", "true")
    only_discount = request.args.get("discount_only", "0") in ("1", "true")

    is_promo_mode = (not raw_query or raw_query.lower() in ("promotions", "акції", "знижки", "акція") or only_promo)
    query = raw_query if raw_query else "🔥 Гарячі акції"

    atb_items = []
    silpo_items = []
    errors = {}

    if is_promo_mode:
        def fetch_promo_atb():
            try:
                return atb_api.get_atb_promos_page(page=page, urls_per_batch=2, timeout=12)
            except Exception as e:
                errors["atb"] = str(e)
                return []

        def fetch_promo_silpo():
            try:
                s_limit = 50 if store == "all" else 100
                s_offset = (page - 1) * s_limit
                return silpo_api.get_silpo_promotions(branch_id=branch_id, limit=s_limit, offset=s_offset, timeout=12)
            except Exception as e:
                errors["silpo"] = str(e)
                return []

        if store == "all":
            with ThreadPoolExecutor(max_workers=2) as executor:
                future_atb = executor.submit(fetch_promo_atb)
                future_silpo = executor.submit(fetch_promo_silpo)
                atb_items = future_atb.result() or []
                silpo_items = future_silpo.result() or []
        elif store == "atb":
            atb_items = fetch_promo_atb()
        elif store == "silpo":
            silpo_items = fetch_promo_silpo()
    else:
        cat_id, cat_info = find_smart_category(raw_query)
        if cat_info:
            subq_batches = cat_info["sub_queries"]
            current_subqueries = subq_batches[(page - 1) % len(subq_batches)]

            def fetch_atb():
                items = []
                seen = set()
                try:
                    with ThreadPoolExecutor(max_workers=len(current_subqueries)) as ex:
                        for res in ex.map(lambda q: atb_api.search_atb(q, timeout=12).get("products", []), current_subqueries):
                            for p in res:
                                k = p.get("product_url") or p.get("title")
                                if k not in seen:
                                    seen.add(k)
                                    items.append(p)
                    return items
                except Exception as e:
                    errors["atb"] = str(e)
                    return []

            def fetch_silpo():
                items = []
                seen = set()
                try:
                    s_limit = 35 if store == "all" else 60
                    with ThreadPoolExecutor(max_workers=len(current_subqueries)) as ex:
                        for res in ex.map(lambda q: silpo_api.search_products(q, branch_id=branch_id, limit=s_limit, only_promo=only_promo, timeout=12).get("items", []), current_subqueries):
                            for p in res:
                                k = p.get("product_url") or p.get("title")
                                if k not in seen:
                                    seen.add(k)
                                    items.append(p)
                    return items
                except Exception as e:
                    errors["silpo"] = str(e)
                    return []
        else:
            def fetch_atb():
                try:
                    res = atb_api.search_atb(raw_query, page=page, timeout=15)
                    return res.get("products", [])
                except Exception as e:
                    errors["atb"] = str(e)
                    return []

            def fetch_silpo():
                try:
                    if store == "silpo":
                        s_limit = 100
                        s_offset = (page - 1) * 100
                    else:
                        s_limit = 76
                        s_offset = (page - 1) * 76
                    res = silpo_api.search_products(
                        query=raw_query,
                        branch_id=branch_id,
                        limit=s_limit,
                        offset=s_offset,
                        only_promo=only_promo,
                        timeout=15
                    )
                    return res.get("items", [])
                except Exception as e:
                    errors["silpo"] = str(e)
                    return []

        if store == "all":
            with ThreadPoolExecutor(max_workers=2) as executor:
                future_atb = executor.submit(fetch_atb)
                future_silpo = executor.submit(fetch_silpo)
                atb_items = future_atb.result() or []
                silpo_items = future_silpo.result() or []
        elif store == "atb":
            atb_items = fetch_atb()
        elif store == "silpo":
            silpo_items = fetch_silpo()

    if store == "all":
        # Чергуємо товари з АТБ та Сільпо для зручного порівняння поруч
        combined = []
        max_len = max(len(atb_items), len(silpo_items))
        for i in range(max_len):
            if i < len(atb_items):
                combined.append(atb_items[i])
            if i < len(silpo_items):
                combined.append(silpo_items[i])
        products = combined
    elif store == "atb":
        products = atb_items
    elif store == "silpo":
        products = silpo_items
    else:
        return jsonify({"success": False, "error": f"Невідомий магазин: {store}"}), 400

    # Фільтр тільки зі знижкою
    if only_discount:
        products = [p for p in products if p.get("has_discount") or (p.get("discount_percent") and p.get("discount_percent") > 0)]

    # Сортування
    if sort == "price_asc":
        products.sort(key=lambda x: (x.get("price") if x.get("price") is not None else 999999))
    elif sort == "price_desc":
        products.sort(key=lambda x: (x.get("price") or 0), reverse=True)
    elif sort == "discount_desc":
        products.sort(key=lambda x: (x.get("discount_percent") or 0), reverse=True)
    elif sort == "saving_desc":
        products.sort(key=lambda x: (x.get("saving") or ((x.get("old_price") - x.get("price")) if (x.get("old_price") and x.get("price") and x.get("old_price") > x.get("price")) else 0)), reverse=True)
    elif sort == "title_asc":
        products.sort(key=lambda x: str(x.get("title", "")).lower())
    elif sort == "title_desc":
        products.sort(key=lambda x: str(x.get("title", "")).lower(), reverse=True)
    elif sort == "unit_price_asc":
        def _get_u_asc(x):
            pr = x.get("price")
            if not pr or pr <= 0:
                return 999999
            text = f"{x.get('title', '')} {x.get('weight', '')}".lower()
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:кг|kg)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) / 10.0
                except Exception:
                    pass
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:л|l)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) / 10.0
                except Exception:
                    pass
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:г|гр|g)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) * 100.0
                except Exception:
                    pass
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:мл|ml)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) * 100.0
                except Exception:
                    pass
            return 999999
        products.sort(key=_get_u_asc)
    elif sort == "unit_price_desc":
        def _get_u_desc(x):
            pr = x.get("price")
            if not pr or pr <= 0:
                return -1
            text = f"{x.get('title', '')} {x.get('weight', '')}".lower()
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:кг|kg)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) / 10.0
                except Exception:
                    pass
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:л|l)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) / 10.0
                except Exception:
                    pass
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:г|гр|g)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) * 100.0
                except Exception:
                    pass
            m = re.search(r'(\d+[\.,]?\d*)\s*(?:мл|ml)(?![а-яa-z0-9])', text)
            if m:
                try:
                    return (pr / float(m.group(1).replace(',', '.'))) * 100.0
                except Exception:
                    pass
            return -1
        products.sort(key=_get_u_desc, reverse=True)

    # Гарантуємо наявність як image, так і image_url для коректного відображення фото
    for p in products:
        img = p.get("image") or p.get("image_url") or p.get("image_large_url") or ""
        p["image"] = img
        p["image_url"] = img
        if not p.get("image_large_url"):
            p["image_large_url"] = img

    return jsonify({
        "success": True,
        "query": query,
        "store": store,
        "page": page,
        "limit": limit,
        "is_promo": is_promo_mode,
        "total": len(products),
        "atb_count": len(atb_items),
        "silpo_count": len(silpo_items),
        "items": products,
        "has_more": len(products) > 0,
        "errors": errors if errors else None
    })


@app.route("/api/catalog/page/<int:page>")
def api_catalog_page(page):
    """Каталог акцій АТБ (сумісність)."""
    try:
        res = atb_api.get_atb_catalog(page=page)
        return jsonify({
            "products": res.get("products", []),
            "page": page,
            "count": len(res.get("products", []))
        })
    except Exception as e:
        return jsonify({"products": [], "error": str(e)}), 200


@app.route("/api/search/page")
def api_search_page():
    """Посторінковий пошук АТБ."""
    query = request.args.get("q", "").strip()
    page = int(request.args.get("page", 1))
    if not query:
        return jsonify({"products": [], "total_pages": 1})

    try:
        res = atb_api.search_atb(query, page=page)
        return jsonify({
            "products": res.get("products", []),
            "page": page,
            "total_pages": 1,
            "count": len(res.get("products", []))
        })
    except Exception as e:
        return jsonify({"products": [], "error": str(e)}), 200


@app.route("/api/details")
def api_details():
    """
    Уніфікований ендпоінт для 3D Flip Card:
    Повертає словник характеристик для картки товару як АТБ, так і Сільпо.
    """
    url = request.args.get("url", "").strip()
    slug = request.args.get("slug", "").strip()
    store = request.args.get("store", "").strip().lower()
    branch_id = request.args.get("branch_id", silpo_api.DEFAULT_BRANCH_ID)

    if "atbmarket.com" in url or store == "atb":
        chars = atb_api.get_atb_details(url)
        return jsonify(chars)

    if not slug and "silpo.ua/product/" in url:
        slug = url.split("/product/")[-1].split("?")[0].strip("/")

    if slug or store == "silpo":
        try:
            details = silpo_api.get_product_details(slug, branch_id=branch_id)
            chars = details.get("characteristics", {})
            if details.get("country") and "Країна" not in chars:
                chars["Країна"] = details["country"]
            if details.get("brand") and "Бренд" not in chars:
                chars["Бренд"] = details["brand"]
            if details.get("ingredients") and "Склад" not in chars:
                chars["Склад"] = details["ingredients"]
            nutr = details.get("nutrition", {})
            if nutr.get("calories") and "Калорійність" not in chars:
                chars["Калорійність (ккал)"] = str(nutr["calories"])
            if nutr.get("protein") and "Білки" not in chars:
                chars["Білки (г)"] = str(nutr["protein"])
            if nutr.get("fat") and "Жири" not in chars:
                chars["Жири (г)"] = str(nutr["fat"])
            if nutr.get("carbohydrates") and "Вуглеводи" not in chars:
                chars["Вуглеводи (г)"] = str(nutr["carbohydrates"])

            return jsonify(chars)
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    return jsonify({"error": "Не вказано URL або slug товару"}), 400


# ═══════════════════════════════════════════
#   СПИСОК БАЖАНОГО (WISHLIST)
# ═══════════════════════════════════════════
@app.route("/api/wishlist", methods=["GET", "POST", "DELETE"])
def api_wishlist():
    db = load_db()
    wishlist = db.get("wishlist", [])

    if request.method == "GET":
        return jsonify({"success": True, "wishlist": wishlist, "count": len(wishlist)})

    elif request.method == "POST":
        item = request.get_json() or {}
        item_id = str(item.get("id") or item.get("link") or item.get("title"))
        if not any(str(w.get("id") or w.get("link") or w.get("title")) == item_id for w in wishlist):
            item["saved_at"] = time.time()
            wishlist.insert(0, item)
            db["wishlist"] = wishlist
            save_db(db)
        return jsonify({"success": True, "wishlist": wishlist, "count": len(wishlist)})

    elif request.method == "DELETE":
        item_id = request.args.get("id")
        if item_id:
            wishlist = [w for w in wishlist if str(w.get("id") or w.get("link") or w.get("title")) != str(item_id)]
        else:
            wishlist = []
        db["wishlist"] = wishlist
        save_db(db)
        return jsonify({"success": True, "wishlist": wishlist, "count": len(wishlist)})


# ═══════════════════════════════════════════
#   РОЗУМНИЙ СПЛІТ-КОШИК & ОПТИМІЗАТОР
# ═══════════════════════════════════════════
@app.route("/api/basket", methods=["GET", "POST", "DELETE"])
def api_basket():
    db = load_db()
    basket = db.get("basket", [])

    if request.method == "GET":
        return jsonify({"success": True, "basket": basket, "count": len(basket)})

    elif request.method == "POST":
        item = request.get_json() or {}
        item_id = str(item.get("id") or item.get("link") or item.get("title"))
        qty = int(item.get("quantity", 1))

        found = False
        for b in basket:
            if str(b.get("id") or b.get("link") or b.get("title")) == item_id:
                b["quantity"] = max(1, b.get("quantity", 1) + qty) if "set_exact_qty" not in item else max(1, qty)
                found = True
                break
        if not found:
            item["quantity"] = max(1, qty)
            basket.append(item)

        db["basket"] = basket
        save_db(db)
        return jsonify({"success": True, "basket": basket, "count": len(basket)})

    elif request.method == "DELETE":
        item_id = request.args.get("id")
        if item_id:
            basket = [b for b in basket if str(b.get("id") or b.get("link") or b.get("title")) != str(item_id)]
        else:
            basket = []
        db["basket"] = basket
        save_db(db)
        return jsonify({"success": True, "basket": basket, "count": len(basket)})


def clean_title_for_search(title: str) -> str:
    """Вилучає грамування та зайві службові слова для пошуку аналога."""
    t = re.sub(r"\d+[\.,]?\d*\s*(?:кг|г|л|мл|шт|%)\b", "", title, flags=re.IGNORECASE)
    t = re.sub(r"[«»\"'()\[\]]", " ", t)
    words = [w for w in t.split() if len(w) > 2 and w.lower() not in ("товар", "акція", "пакет", "плівка", "тетра")]
    return " ".join(words[:3]) or title[:15]


@app.route("/api/basket/optimize", methods=["POST", "GET"])
def optimize_basket():
    """
    Калькулятор максимальної вигоди (Спліт-кошик):
    Порівнює вартість кошика:
    - якщо купити все в АТБ
    - якщо купити все в Сільпо
    - якщо взяти кожну позицію там, де вона дешевша (Спліт)
    """
    db = load_db()
    req_data = request.get_json() if request.method == "POST" else None
    basket = (req_data.get("basket") if req_data else None) or db.get("basket", [])

    if not basket:
        return jsonify({
            "success": True,
            "total_atb": 0,
            "total_silpo": 0,
            "total_split": 0,
            "savings": 0,
            "items": []
        })

    optimized_items = []
    total_atb = 0.0
    total_silpo = 0.0
    total_split = 0.0

    def find_peer_price(item):
        orig_store = item.get("store", "atb")
        title = item.get("title", "")
        clean_q = clean_title_for_search(title)
        peer_store = "silpo" if orig_store == "atb" else "atb"

        peer_price = None
        peer_title = None
        peer_url = None

        try:
            if peer_store == "silpo":
                res = silpo_api.search_products(clean_q, limit=3)
                p_items = res.get("items", [])
                if p_items:
                    best = min(p_items, key=lambda x: x.get("price") or 9999)
                    peer_price = best.get("price")
                    peer_title = best.get("title")
                    peer_url = best.get("product_url")
            else:
                res = atb_api.search_atb(clean_q, page=1)
                p_items = res.get("products", [])
                if p_items:
                    best = min(p_items, key=lambda x: x.get("price") or 9999)
                    peer_price = best.get("price")
                    peer_title = best.get("title")
                    peer_url = best.get("product_url") or best.get("link")
        except Exception:
            pass

        return peer_store, peer_price, peer_title, peer_url

    with ThreadPoolExecutor(max_workers=4) as executor:
        peers = list(executor.map(find_peer_price, basket))

    for item, (peer_store, peer_price, peer_title, peer_url) in zip(basket, peers):
        qty = int(item.get("quantity", 1))
        curr_p = float(item.get("price") or 0.0)
        curr_store = item.get("store", "atb")

        if curr_store == "atb":
            p_atb = curr_p
            p_silpo = float(peer_price) if peer_price is not None else round(curr_p * 1.05, 2)
        else:
            p_silpo = curr_p
            p_atb = float(peer_price) if peer_price is not None else round(curr_p * 0.95, 2)

        # Найкращий магазин для цієї позиції
        if p_atb <= p_silpo:
            rec_store = "atb"
            rec_store_name = "АТБ"
            rec_price = p_atb
            diff = round(p_silpo - p_atb, 2)
        else:
            rec_store = "silpo"
            rec_store_name = "Сільпо"
            rec_price = p_silpo
            diff = round(p_atb - p_silpo, 2)

        total_atb += p_atb * qty
        total_silpo += p_silpo * qty
        total_split += rec_price * qty

        optimized_items.append({
            "id": item.get("id"),
            "title": item.get("title"),
            "quantity": qty,
            "unit": item.get("unit", "шт"),
            "weight": item.get("weight", ""),
            "price_atb": round(p_atb, 2),
            "price_silpo": round(p_silpo, 2),
            "recommended_store": rec_store,
            "recommended_store_name": rec_store_name,
            "item_saving": round(diff * qty, 2),
            "peer_title": peer_title,
            "peer_url": peer_url,
            "image_url": item.get("image_url")
        })

    min_single = min(total_atb, total_silpo)
    savings = max(0.0, round(min_single - total_split, 2))

    return jsonify({
        "success": True,
        "total_atb": round(total_atb, 2),
        "total_silpo": round(total_silpo, 2),
        "total_split": round(total_split, 2),
        "savings": savings,
        "items": optimized_items
    })


# ═══════════════════════════════════════════
#   РАДАР РЕКОРДНИХ ЗНИЖОК ДНЯ (CRAZY DEALS)
# ═══════════════════════════════════════════
@app.route("/api/crazy-deals")
def crazy_deals():
    """
    Повертає товари з найглибшими знижками (від 30-50%) та максимальною економією
    одночасно з АТБ та з Сільпо.
    """
    deals = []

    def load_atb_deals():
        try:
            res = atb_api.get_atb_catalog(page=1)
            return [p for p in res.get("products", []) if (p.get("discount_percent") or 0) >= 30 or (p.get("saving") or 0) >= 20]
        except Exception:
            return []

    def load_silpo_deals():
        try:
            items = silpo_api.get_silpo_promotions(limit=36)
            return [p for p in items if (p.get("discount_percent") or 0) >= 30 or (p.get("saving") or 0) >= 20]
        except Exception:
            return []

    with ThreadPoolExecutor(max_workers=2) as executor:
        f_atb = executor.submit(load_atb_deals)
        f_silpo = executor.submit(load_silpo_deals)
        atb_deals = f_atb.result()
        silpo_deals = f_silpo.result()

    # Змішуємо товари з обох мереж для рівномірного представлення
    max_len = max(len(atb_deals), len(silpo_deals))
    for i in range(max_len):
        if i < len(atb_deals):
            deals.append(atb_deals[i])
        if i < len(silpo_deals):
            deals.append(silpo_deals[i])

    deals.sort(key=lambda x: (x.get("discount_percent") or 0, x.get("saving") or 0), reverse=True)

    return jsonify({
        "success": True,
        "deals": deals[:40],
        "count": len(deals[:40])
    })


# ═══════════════════════════════════════════
#   АНАЛІЗАТОР ЧЕКІВ (RECEIPT INSPECTOR) + AI СКАНЕР ФОТО
# ═══════════════════════════════════════════
@app.route("/api/receipt/ai-scan", methods=["POST"])
def receipt_ai_scan():
    """
    Розпізнає сфотографований чек через Google Gemini Vision AI (модель gemini-flash-lite-latest)
    та порівнює кожен товар між АТБ та Сільпо.
    """
    import base64
    ai_token = os.environ.get("AI_STUDIO_TOKEN", "").strip()
    img_b64 = None
    mime_type = "image/jpeg"

    # Якщо файл завантажено через multipart/form-data
    if "photo" in request.files:
        f = request.files["photo"]
        mime_type = f.mimetype or "image/jpeg"
        img_b64 = base64.b64encode(f.read()).decode("utf-8")
    else:
        req_json = request.get_json() or {}
        img_b64 = req_json.get("image_base64")
        mime_type = req_json.get("mime_type", "image/jpeg")

    if not img_b64:
        return jsonify({"success": False, "error": "Не передано зображення чека"}), 400

    prompt = """Ти — професійний аналізатор чеків супермаркетів України.
Поглянь на цей чек і розпізнай усі куплені харчові та господарські товари та їхні фінальні ціни (в грн).
Ігноруй пакети-майки, здачу, дисконтні картки, податки та підсумки.
Поверни ВИНЯТКОВО валідний JSON масив без жодного іншого тексту, у форматі:
[
  {"name": "Молоко Селянське 2.5%", "price": 38.50},
  {"name": "Масло Яготинське 73%", "price": 79.90}
]"""

    try:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-flash-lite-latest:generateContent?key={ai_token}"
        payload = {
            "contents": [{
                "parts": [
                    {"text": prompt},
                    {"inlineData": {"mimeType": mime_type, "data": img_b64}}
                ]
            }]
        }
        req = urllib.request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req, timeout=25) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            raw_ai = data["candidates"][0]["content"]["parts"][0]["text"].strip()

        # Витягуємо JSON із відповіді
        clean_json = raw_ai
        if "```json" in clean_json:
            clean_json = clean_json.split("```json")[1].split("```")[0].strip()
        elif "```" in clean_json:
            clean_json = clean_json.split("```")[1].split("```")[0].strip()

        items = json.loads(clean_json)
    except Exception as e:
        return jsonify({"success": False, "error": f"Помилка ШІ розпізнавання: {e}"}), 500

    # Передаємо розпізнані товари до аналізатора чеків
    return process_receipt_items(items)


@app.route("/api/receipt/analyze", methods=["POST"])
def analyze_receipt():
    """Приймає текст або масив товарів чека і порівнює ціни в АТБ та Сільпо."""
    data = request.get_json() or {}
    raw_text = data.get("text", "")
    items_input = data.get("items", [])

    lines = []
    if raw_text:
        for line in raw_text.split("\n"):
            line = line.strip()
            if not line:
                continue
            m = re.search(r"^(.*?)\s+([\d\.,]+)\s*(?:грн|₴)?$", line, re.IGNORECASE)
            if m:
                lines.append({"name": m.group(1).strip(), "price": float(m.group(2).replace(",", "."))})
            else:
                lines.append({"name": line, "price": None})
    elif items_input:
        lines = items_input

    if not lines:
        return jsonify({"success": False, "error": "Не знайдено товарів у чеку"}), 400

    return process_receipt_items(lines)


def process_receipt_items(lines):
    results = []
    total_receipt = 0.0
    total_atb = 0.0
    total_silpo = 0.0

    def match_item(it):
        name = it.get("name", "")
        p_orig = it.get("price")
        clean_q = clean_title_for_search(name)

        p_atb = None
        p_silpo = None
        match_atb = None
        match_silpo = None

        try:
            r_atb = atb_api.search_atb(clean_q, page=1)
            if r_atb.get("products"):
                best_a = r_atb["products"][0]
                p_atb = best_a.get("price")
                match_atb = best_a.get("title")
        except Exception:
            pass

        try:
            r_silpo = silpo_api.search_products(clean_q, limit=1)
            if r_silpo.get("items"):
                best_s = r_silpo["items"][0]
                p_silpo = best_s.get("price")
                match_silpo = best_s.get("title")
        except Exception:
            pass

        return {
            "name": name,
            "original_price": p_orig,
            "atb_price": p_atb,
            "atb_title": match_atb,
            "silpo_price": p_silpo,
            "silpo_title": match_silpo
        }

    with ThreadPoolExecutor(max_workers=4) as executor:
        results = list(executor.map(match_item, lines[:15]))

    for r in results:
        if r["original_price"]:
            total_receipt += r["original_price"]
        if r["atb_price"]:
            total_atb += r["atb_price"]
        elif r["original_price"]:
            total_atb += r["original_price"]

        if r["silpo_price"]:
            total_silpo += r["silpo_price"]
        elif r["original_price"]:
            total_silpo += r["original_price"]

    return jsonify({
        "success": True,
        "total_receipt": round(total_receipt, 2),
        "total_atb": round(total_atb, 2),
        "total_silpo": round(total_silpo, 2),
        "diff_atb": round(total_receipt - total_atb, 2) if total_receipt else 0,
        "diff_silpo": round(total_receipt - total_silpo, 2) if total_receipt else 0,
        "items": results
    })


# ═══════════════════════════════════════════
#   ПОШУК ЗА ШТРИХКОДОМ (BARCODE LOOKUP)
# ═══════════════════════════════════════════
@app.route("/api/barcode/<code_or_query>")
def barcode_lookup(code_or_query):
    """
    Пошук товару за числовим штрихкодом (EAN-13) або назвою.
    Використовує пошук за штрихкодом Сільпо (searchV2), Open Food Facts та зіставлення в АТБ.
    """
    code = code_or_query.strip()
    silpo_items = []
    atb_items = []
    resolved_title = ""

    # 1. Пошук за штрихкодом у Сільпо (через searchV2)
    try:
        silpo_items = silpo_api.search_silpo_barcode(code)
    except Exception:
        silpo_items = []

    if silpo_items:
        resolved_title = silpo_items[0].get("title", "")
    else:
        # 2. Фолбек на базу Open Food Facts
        try:
            off_url = f"https://world.openfoodfacts.org/api/v0/product/{code}.json"
            req_off = urllib.request.Request(off_url, headers={"User-Agent": "ATB-Silpo/1.0"})
            with urllib.request.urlopen(req_off, timeout=3) as resp:
                off_data = json.loads(resp.read().decode())
                if off_data.get("status") == 1:
                    p = off_data.get("product", {})
                    resolved_title = p.get("product_name_uk") or p.get("product_name") or p.get("brands", "")
        except Exception:
            pass

    # 3. Шукаємо відповідник в АТБ за назвою товару
    if resolved_title:
        clean_q = clean_title_for_search(resolved_title)
        try:
            atb_res = atb_api.search_atb(clean_q)
            atb_items = atb_res.get("products", [])
        except Exception:
            pass
        if not silpo_items:
            try:
                silpo_items = silpo_api.search_products(clean_q, limit=4).get("items", [])
            except Exception:
                pass
    else:
        # Прямий текстовий пошук у двох мережах
        with ThreadPoolExecutor(max_workers=2) as executor:
            f_a = executor.submit(atb_api.search_atb, code)
            f_s = executor.submit(silpo_api.search_products, code, limit=4)
            atb_items = f_a.result().get("products", [])
            silpo_items = f_s.result().get("items", [])

    items = silpo_items + atb_items
    return jsonify({
        "success": True,
        "barcode": code,
        "resolved_title": resolved_title,
        "total": len(items),
        "items": items
    })


# ═══════════════════════════════════════════
#   НАЛАШТУВАННЯ КАРТОК ТА КЕШБЕКУ
# ═══════════════════════════════════════════
@app.route("/api/settings", methods=["GET", "POST"])
def api_settings():
    db = load_db()
    if request.method == "POST":
        new_settings = request.get_json() or {}
        db["settings"] = {**db.get("settings", {}), **new_settings}
        save_db(db)
        return jsonify({"success": True, "settings": db["settings"]})
    return jsonify({"success": True, "settings": db.get("settings", {})})


# ═══════════════════════════════════════════
#   ЕКСПОРТ ДАНИХ (EXCEL / CSV / JSON)
# ═══════════════════════════════════════════
@app.route("/api/export", methods=["POST"])
def export_products():
    """Експорт обраних товарів (АТБ + Сільпо) в Excel, CSV або JSON."""
    data = request.get_json() or {}
    items = data.get("items", [])
    fmt = data.get("format", "xlsx").lower()
    filename = data.get("filename", "supermarket_prices").strip() or "supermarket_prices"

    if not items:
        return jsonify({"success": False, "error": "Немає товарів для експорту"}), 400

    if fmt == "json":
        buf = io.BytesIO(json.dumps(items, ensure_ascii=False, indent=2).encode("utf-8"))
        return send_file(buf, mimetype="application/json", as_attachment=True, download_name=f"{filename}.json")

    elif fmt == "csv":
        buf = io.StringIO()
        fieldnames = [
            "store_name", "title", "price", "old_price", "discount_percent",
            "saving", "is_special_card", "special_card_name", "weight", "product_url"
        ]
        writer = csv.DictWriter(buf, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for it in items:
            writer.writerow(it)

        bytes_io = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
        return send_file(bytes_io, mimetype="text/csv", as_attachment=True, download_name=f"{filename}.csv")

    elif fmt == "xlsx":
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            return export_products_csv(items, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "АТБ та Сільпо Порівняння"

        headers = [
            "Супермаркет", "Назва товару", "Ціна (₴)", "Стара ціна (₴)",
            "Знижка (%)", "Економія (₴)", "Спец. картка", "Фасування / Вага", "Посилання"
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        atb_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
        silpo_fill = PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")

        for row_idx, p in enumerate(items, start=2):
            store_name = p.get("store_name", "АТБ" if p.get("store") == "atb" else "Сільпо")
            disc = f"-{p.get('discount_percent')}%" if p.get("discount_percent") else ""
            card_info = p.get("special_card_name") or ("Так" if p.get("is_special_card") else "")

            ws.append([
                store_name,
                p.get("title", ""),
                p.get("price", 0.0),
                p.get("old_price") or "",
                disc,
                p.get("saving") or "",
                card_info,
                p.get("weight", ""),
                p.get("product_url", "") or p.get("link", "")
            ])

            row_fill = atb_fill if p.get("store") == "atb" else silpo_fill
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = row_fill

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        return send_file(
            excel_buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"{filename}.xlsx"
        )

    return jsonify({"success": False, "error": "Невідомий формат"}), 400


def export_products_csv(items, filename):
    buf = io.StringIO()
    fieldnames = ["store_name", "title", "price", "old_price", "discount_percent", "saving", "product_url"]
    writer = csv.DictWriter(buf, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for it in items:
        writer.writerow(it)
    bytes_io = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
    return send_file(bytes_io, mimetype="text/csv", as_attachment=True, download_name=f"{filename}.csv")


# ═══════════════════════════════════════════
#   SILPO MCP ENDPOINTS (ЗБЕРЕЖЕННЯ ПРОФІЛЮ)
# ═══════════════════════════════════════════
@app.route("/api/mcp/status")
def mcp_status():
    mcp_client._load_tokens()
    status = mcp_client.get_status()

    user_info = None
    if status.get("authorized"):
        try:
            p_res = mcp_client.call_mcp_rpc("tools/call", {"name": "silpo_get_my_profile", "arguments": {}})
            l_res = mcp_client.call_mcp_rpc("tools/call", {"name": "silpo_get_loyalty_info", "arguments": {}})
            prof = p_res.get("result", {}).get("structuredContent", {}).get("profile")
            loyalty = l_res.get("result", {}).get("structuredContent", {}).get("loyalty")
            user_info = {
                "name": f"{prof.get('firstName', '')} {prof.get('lastName', '')}".strip() if prof else None,
                "phone": prof.get("phone") if prof else None,
                "bonus": loyalty.get("balance", {}).get("total") if loyalty else 0,
                "card": loyalty.get("card", {}).get("barcode") if loyalty else None
            }
        except Exception as e:
            user_info = {"error": str(e)}

    status["user"] = user_info
    return jsonify(status)


@app.route("/api/mcp/login")
def mcp_login():
    try:
        auth_url, verifier, state = mcp_client.get_authorization_url()
        _oauth_verifier_cache[state] = verifier
        return redirect(auth_url)
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/mcp/logout")
def mcp_logout():
    mcp_client.access_token = None
    mcp_client.refresh_token = None
    if os.path.exists(TOKEN_FILE):
        try:
            os.remove(TOKEN_FILE)
        except Exception:
            pass
    return redirect("/")


@app.route("/oauth/callback")
def oauth_callback():
    code = request.args.get("code")
    state = request.args.get("state")
    error = request.args.get("error")

    if error:
        return f"<h3>Помилка авторизації Сільпо: {error}</h3><p><a href='/'>Повернутись</a></p>"

    if not code or not state:
        return "<h3>Невірні параметри відповіді OAuth</h3>", 400

    verifier = _oauth_verifier_cache.pop(state, None)
    if not verifier:
        return "<h3>Сесію авторизації не знайдено або термін дії вичерпано.</h3>", 400

    try:
        mcp_client.exchange_code_for_token(code, verifier)
        return redirect("/?mcp_authorized=1")
    except Exception as e:
        return f"<h3>Не вдалося отримати токен MCP: {e}</h3><p><a href='/'>Повернутись</a></p>"




# ═══════════════════════════════════════════
#   МОДУЛЬ АНАЛІТИКИ ТА СТАТИСТИКИ АДМІНА (7740)
#   (Підтримка Vercel KV / Upstash Redis + локальний fallback)
# ═══════════════════════════════════════════
STATS_FILE = os.path.join(BASE_DIR, "stats.json")
TMP_STATS_FILE = "/tmp/stats.json"
_stats_lock = threading.RLock()

REDIS_KV_KEY = "loot_analytics_stats_v1"

def get_kv_credentials():
    url = (os.environ.get("KV_REST_API_URL") or 
           os.environ.get("UPSTASH_REDIS_REST_URL") or "").strip()
    token = (os.environ.get("KV_REST_API_TOKEN") or 
             os.environ.get("UPSTASH_REDIS_REST_TOKEN") or "").strip()
    if url and token:
        return url.rstrip("/"), token
    return None, None

def kv_redis_get(key):
    url, token = get_kv_credentials()
    if not url or not token:
        return None
    try:
        req_url = f"{url}/get/{key}"
        req = urllib.request.Request(req_url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req, timeout=4) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            raw = data.get("result")
            if raw:
                if isinstance(raw, str):
                    return json.loads(raw)
                return raw
    except Exception as e:
        print(f"Vercel KV Read error ({key}): {e}")
    return None

def kv_redis_set(key, data_obj):
    url, token = get_kv_credentials()
    if not url or not token:
        return False
    try:
        # Використовуємо універсальний POST до Upstash REST API
        val_str = json.dumps(data_obj, ensure_ascii=False)
        cmd_payload = json.dumps(["SET", key, val_str]).encode("utf-8")
        req = urllib.request.Request(
            url,
            data=cmd_payload,
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": "application/json"
            },
            method="POST"
        )
        with urllib.request.urlopen(req, timeout=4) as resp:
            return resp.status in (200, 201)
    except Exception as e:
        print(f"Vercel KV Write error ({key}): {e}")
        return False

def get_stats_file_path():
    try:
        test_file = os.path.join(BASE_DIR, ".write_test")
        with open(test_file, "w") as f:
            f.write("1")
        os.remove(test_file)
        return STATS_FILE
    except Exception:
        return TMP_STATS_FILE

def get_default_stats():
    return {
        "summary": {
            "total_visits": 0,
            "unique_visitors": 0,
            "total_searches": 0,
            "total_clicks": 0
        },
        "visitors": {},
        "search_queries": {},
        "recent_events": []
    }

def load_stats():
    with _stats_lock:
        # 1. Спроба завантажити з Vercel KV / Upstash Redis
        kv_data = kv_redis_get(REDIS_KV_KEY)
        if kv_data and isinstance(kv_data, dict) and "summary" in kv_data:
            return kv_data

        # 2. Локальний файл fallback
        file_path = get_stats_file_path()
        if not os.path.exists(file_path):
            stats = get_default_stats()
            save_stats(stats)
            return stats
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print(f"Помилка читання stats.json: {e}")
            return get_default_stats()

def save_stats(data):
    with _stats_lock:
        # 1. Запис у Vercel KV / Upstash Redis
        kv_success = kv_redis_set(REDIS_KV_KEY, data)

        # 2. Одночасно зберігаємо у локальний файл/tmp для бекапу
        file_path = get_stats_file_path()
        try:
            with open(file_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            print(f"Помилка запису в {file_path}: {e}")
            return kv_success

@app.route("/api/analytics/track", methods=["POST"])
def track_analytics():
    try:
        payload = request.get_json(silent=True) or {}
        event_type = payload.get("type", "visit") # visit, search, click, product_view
        raw_vid = str(payload.get("visitor_id") or "anon").strip()
        device_fp = str(payload.get("device_fingerprint") or "").strip()
        data_text = str(payload.get("data") or "").strip()
        device_info = payload.get("device") or {}

        # Час за Києвом (+03:00) / UTC
        current_time = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())

        # Визначення реальної IP адреси
        forwarded = request.headers.get("X-Forwarded-For")
        if forwarded:
            ip = forwarded.split(",")[0].strip()
        else:
            ip = request.remote_addr or "127.0.0.1"

        # Детектор автоматичних роботів/ботів (наприклад, Vercel preview crawler, Googlebot, Lighthouse)
        ua_raw = request.headers.get("User-Agent", "")
        ua_lower = ua_raw.lower()
        gpu_lower = str(device_info.get("gpu") or "").lower()
        res_screen = str(device_info.get("screenResolution") or "")

        is_bot = False
        if ("swiftshader" in gpu_lower or "headless" in ua_lower or "bot" in ua_lower or 
            "crawl" in ua_lower or "spider" in ua_lower or "lighthouse" in ua_lower or 
            (res_screen == "800×600" and "google" in gpu_lower) or
            ip.startswith("13.57.") or ip.startswith("54.") or ip.startswith("34.")):
            is_bot = True

        stats = load_stats()
        summary = stats.setdefault("summary", {
            "total_visits": 0,
            "unique_visitors": 0,
            "total_searches": 0,
            "total_clicks": 0
        })
        visitors = stats.setdefault("visitors", {})
        search_queries = stats.setdefault("search_queries", {})
        recent_events = stats.setdefault("recent_events", [])

        # ═══════════════════════════════════════════
        #   РОЗУМНЕ СКЛЕЮВАННЯ ПРИСТРОЇВ (SMART FINGERPRINT MATCHING)
        # ═══════════════════════════════════════════
        target_vid = raw_vid

        if target_vid not in visitors and not is_bot:
            # 1. Пошук за цифровим відбитком заліза (Device Fingerprint)
            matched_id = None
            if device_fp:
                for vid, vdata in visitors.items():
                    if vdata.get("device_fingerprint") == device_fp:
                        matched_id = vid
                        break

            # 2. Якщо fingerprint не знайдено, перевіряємо ідентичність параметрів заліза + підмережі IP
            if not matched_id and device_info:
                cur_screen = device_info.get("screenResolution")
                cur_gpu = device_info.get("gpu")
                cur_model = device_info.get("model")
                cur_tz = device_info.get("timezone")

                ip_parts = ip.split(".")
                ip_subnet = ".".join(ip_parts[:3]) if len(ip_parts) == 4 else ip

                for vid, vdata in visitors.items():
                    vdev = vdata.get("device") or {}
                    vip = vdata.get("ip") or ""
                    vip_subnet = ".".join(vip.split(".")[:3]) if len(vip.split(".")) == 4 else vip

                    # Суворе порівняння: той самий екран + відеочіп + часовий пояс + підмережа IP/модель
                    if (cur_screen and cur_screen == vdev.get("screenResolution") and
                        cur_gpu and cur_gpu == vdev.get("gpu") and
                        cur_tz and cur_tz == vdev.get("timezone") and
                        (ip_subnet == vip_subnet or cur_model == vdev.get("model"))):
                        matched_id = vid
                        break

            if matched_id:
                target_vid = matched_id

        # Ініціалізація нового або оновлення існуючого
        if target_vid not in visitors:
            visitors[target_vid] = {
                "first_seen": current_time,
                "last_seen": current_time,
                "visit_count": 0,
                "ip": ip,
                "device_fingerprint": device_fp,
                "is_bot": is_bot,
                "device": device_info,
                "searches": [],
                "clicks_count": 0
            }

        v_record = visitors[target_vid]
        v_record["last_seen"] = current_time
        v_record["ip"] = ip
        if device_fp:
            v_record["device_fingerprint"] = device_fp
        if is_bot:
            v_record["is_bot"] = True

        if device_info:
            v_record["device"] = device_info

        # Зберігаємо сесійну статистику
        session_info = payload.get("session") or {}
        v_record["products_viewed"] = v_record.get("products_viewed", 0) + session_info.get("products_viewed", 0)
        if session_info.get("duration_sec"):
            v_record["total_session_sec"] = v_record.get("total_session_sec", 0) + session_info["duration_sec"]

        # Обробка конкретної події
        if event_type == "visit":
            summary["total_visits"] = summary.get("total_visits", 0) + 1
            v_record["visit_count"] = v_record.get("visit_count", 0) + 1
            event_desc = f"Відкрив(ла) сайт"

        elif event_type == "search":
            if data_text and data_text != "7740":
                summary["total_searches"] = summary.get("total_searches", 0) + 1
                clean_q = data_text.strip().lower()
                search_queries[clean_q] = search_queries.get(clean_q, 0) + 1

                # Додаємо до історії пошуків цього користувача
                u_searches = v_record.setdefault("searches", [])
                if not u_searches or u_searches[-1] != clean_q:
                    u_searches.append(clean_q)
                    if len(u_searches) > 50:
                        v_record["searches"] = u_searches[-50:]
                event_desc = f"Пошук: «{data_text}»"
            else:
                event_desc = "Пошук"

        elif event_type == "click":
            summary["total_clicks"] = summary.get("total_clicks", 0) + 1
            v_record["clicks_count"] = v_record.get("clicks_count", 0) + 1
            event_desc = f"Клік: {data_text}" if data_text else "Клік по елементу"

        elif event_type == "product_view":
            v_record["products_viewed"] = v_record.get("products_viewed", 0) + 1
            event_desc = f"Переглянув товар: {data_text}"

        elif event_type == "session_ping":
            event_desc = f"Активна сесія: {data_text}"

        else:
            event_desc = f"{event_type}: {data_text}"

        # Оновлення кількості унікальних відвідувачів
        summary["unique_visitors"] = len(visitors)

        # Формування рядка для швидкого опису пристрою
        dev_summary = ""
        if device_info:
            d_mod = device_info.get("model") or device_info.get("type", "")
            d_br = device_info.get("browser", "")
            dev_summary = f"{d_mod} ({d_br})".strip(" •")
        if not dev_summary:
            dev_summary = ua_raw[:50]

        # Додавання події до журналу (останні 300 подій; session_ping не записуємо щоб не спамити)
        if event_type != "session_ping":
            recent_events.insert(0, {
                "time": current_time,
                "visitor_id": target_vid,
                "custom_name": v_record.get("custom_name", ""),
                "type": event_type,
                "details": event_desc,
                "ip": ip,
                "device": dev_summary,
                "is_bot": is_bot
            })
        if len(recent_events) > 300:
            stats["recent_events"] = recent_events[:300]

        save_stats(stats)
        return jsonify({"success": True, "canonical_visitor_id": target_vid})
    except Exception as e:
        print(f"Помилка трекінгу аналітики: {e}")
        return jsonify({"success": False, "error": str(e)}), 500

@app.route("/api/admin/stats", methods=["GET"])
def get_admin_stats():
    pin = request.args.get("pin") or request.headers.get("X-Admin-Pin")
    if pin != "7740":
        return jsonify({"success": False, "error": "Невірний код доступу адміна (7740)"}), 403
    stats = load_stats()
    return jsonify({"success": True, "stats": stats})


@app.route("/api/admin/stats/set_name", methods=["POST"])
def set_visitor_name():
    payload = request.get_json(silent=True) or {}
    pin = payload.get("pin") or request.args.get("pin")
    if pin != "7740":
        return jsonify({"success": False, "error": "Невірний код доступу"}), 403
    visitor_id = str(payload.get("visitor_id") or "").strip()
    custom_name = str(payload.get("custom_name") or "").strip()
    if not visitor_id:
        return jsonify({"success": False, "error": "Не вказано visitor_id"}), 400
    
    stats = load_stats()
    visitors = stats.setdefault("visitors", {})
    if visitor_id in visitors:
        visitors[visitor_id]["custom_name"] = custom_name
        save_stats(stats)
        return jsonify({"success": True, "custom_name": custom_name})
    return jsonify({"success": False, "error": "Користувача не знайдено"}), 404


@app.route("/api/admin/stats/delete_visitor", methods=["POST"])
def delete_visitor():
    payload = request.get_json(silent=True) or {}
    pin = payload.get("pin") or request.args.get("pin")
    if pin != "7740":
        return jsonify({"success": False, "error": "Невірний код доступу"}), 403
    vid = str(payload.get("visitor_id") or "").strip()
    if not vid:
        return jsonify({"success": False, "error": "Не вказано visitor_id"}), 400

    stats = load_stats()
    visitors = stats.setdefault("visitors", {})
    if vid in visitors:
        del visitors[vid]
        # Очистити також із журналу останніх подій
        stats["recent_events"] = [e for e in stats.get("recent_events", []) if e.get("visitor_id") != vid]
        stats["summary"]["unique_visitors"] = len(visitors)
        save_stats(stats)
        return jsonify({"success": True, "message": f"Користувача {vid} видалено"})
    return jsonify({"success": False, "error": "Користувача не знайдено"}), 404

@app.route("/api/admin/stats/merge_visitors", methods=["POST"])
def merge_visitors():
    payload = request.get_json(silent=True) or {}
    pin = payload.get("pin") or request.args.get("pin")
    if pin != "7740":
        return jsonify({"success": False, "error": "Невірний код доступу"}), 403
    source_vid = str(payload.get("source_id") or "").strip()
    target_vid = str(payload.get("target_id") or "").strip()
    if not source_vid or not target_vid or source_vid == target_vid:
        return jsonify({"success": False, "error": "Некоректні ID для об'єднання"}), 400

    stats = load_stats()
    visitors = stats.setdefault("visitors", {})
    if source_vid not in visitors or target_vid not in visitors:
        return jsonify({"success": False, "error": "Одного з користувачів не знайдено"}), 404

    src = visitors[source_vid]
    tgt = visitors[target_vid]

    # Зливаємо візити та час
    tgt["visit_count"] = (tgt.get("visit_count") or 1) + (src.get("visit_count") or 1)
    tgt["total_session_sec"] = (tgt.get("total_session_sec") or 0) + (src.get("total_session_sec") or 0)
    tgt["products_viewed"] = (tgt.get("products_viewed") or 0) + (src.get("products_viewed") or 0)
    tgt["clicks_count"] = (tgt.get("clicks_count") or 0) + (src.get("clicks_count") or 0)

    # Зливаємо історію пошуків (унікальні)
    existing_searches = set(tgt.get("searches", []))
    for sq in src.get("searches", []):
        if sq not in existing_searches:
            tgt.setdefault("searches", []).append(sq)

    # Переприв'язуємо останні події
    for ev in stats.get("recent_events", []):
        if ev.get("visitor_id") == source_vid:
            ev["visitor_id"] = target_vid
            if tgt.get("custom_name"):
                ev["custom_name"] = tgt["custom_name"]

    # Видаляємо дубльований профіль
    del visitors[source_vid]
    stats["summary"]["unique_visitors"] = len(visitors)
    save_stats(stats)
    return jsonify({"success": True, "message": "Користувачів успішно об'єднано"})

@app.route("/api/admin/stats/reset", methods=["POST"])
def reset_admin_stats():
    payload = request.get_json(silent=True) or {}
    pin = payload.get("pin") or request.args.get("pin")
    if pin != "7740":
        return jsonify({"success": False, "error": "Невірний код доступу"}), 403
    stats = get_default_stats()
    save_stats(stats)
    return jsonify({"success": True, "message": "Статистику успішно очищено"})

@app.route("/api/admin/stats/download", methods=["GET"])
def download_admin_stats():
    pin = request.args.get("pin")
    if pin != "7740":
        return "Доступ заборонено", 403
    stats = load_stats()
    buf = io.BytesIO()
    buf.write(json.dumps(stats, ensure_ascii=False, indent=2).encode("utf-8"))
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="stats.json", mimetype="application/json")


def open_browser():
    time.sleep(1.2)
    webbrowser.open("http://localhost:5050")


if __name__ == "__main__":
    print("\n" + "=" * 65)
    print("🚀 Запуск сервера «ATB & Сільпо PRO: Порівняння цін та Акцій»...")
    print("🌐 Додаток доступний за адресою: http://localhost:5050")
    print("=" * 65 + "\n")

    threading.Thread(target=open_browser, daemon=True).start()
    app.run(host="0.0.0.0", port=5050, debug=False)
